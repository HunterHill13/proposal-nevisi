#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_audited_proposal_docx.py - Compiles MEDICAL_PROPOSAL_LUPEOL_NDV_V87_REAL_WORLD.docx
with prominent audit status: DRAFT — EVIDENCE AUDIT FAILED
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

sys.path.insert(0, os.path.abspath("scripts"))
from docx_builder import set_p_rtl, add_r


def build_audited_docx():
    print("Building MEDICAL_PROPOSAL_LUPEOL_NDV_V87_REAL_WORLD.docx...")
    doc = Document()
    
    # Configure margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Prominent Audit Failure Warning Banner
    p_warn = doc.add_paragraph()
    set_p_rtl(p_warn, space_before=6, space_after=12, align="center")
    add_r(p_warn, "══════════════════════════════════════════════════════════════════", size_pt=10, bold=True, color_rgb=(0xB0, 0x00, 0x20))
    
    p_warn_text = doc.add_paragraph()
    set_p_rtl(p_warn_text, space_before=4, space_after=4, align="center")
    add_r(p_warn_text, "هشدار ممیزی شواهد واقعی (REAL-WORLD EVIDENCE AUDIT NOTICE)", size_pt=13, bold=True, color_rgb=(0xB0, 0x00, 0x20))
    
    p_status = doc.add_paragraph()
    set_p_rtl(p_status, space_before=4, space_after=4, align="center")
    add_r(p_status, "وضعیت سند: پیش‌نویس مشروط — اعتبارسنجی تجربی ناموفق (DRAFT — EVIDENCE AUDIT FAILED)", size_pt=12, bold=True, color_rgb=(0xC0, 0x00, 0x00))
    
    p_detail = doc.add_paragraph()
    set_p_rtl(p_detail, space_before=4, space_after=8, align="both")
    add_r(p_detail, "بر اساس الزامات ممیزی تجربی نسخه v8.7، این پیش‌نویس به دلیل شناسایی ۴ خطای بحرانی ناهمخوانی مداخله (Refs 8, 18, 23, 24 حاوی اروسین، اوژنول، ماترین و ژولکینولید به جای لوپئول/NDV) و نشت توکن‌های قالبی در ۱۵ پاراگراف پیشینه، فاقد تأییدیه علمی نهایی است و صرفاً به عنوان سند ممیزی خصمانه صادر می‌گردد.", size_pt=10, italic=True, color_rgb=(0x55, 0x55, 0x55))

    p_warn2 = doc.add_paragraph()
    set_p_rtl(p_warn2, space_before=2, space_after=18, align="center")
    add_r(p_warn2, "══════════════════════════════════════════════════════════════════", size_pt=10, bold=True, color_rgb=(0xB0, 0x00, 0x20))

    # Read markdown proposal text
    with open("MEDICAL_PROPOSAL_LUPEOL_NDV.md", "r", encoding="utf-8") as f:
        md_text = f.read()

    lines = md_text.split("\n")
    for line in lines:
        line_s = line.strip()
        if not line_s:
            continue
        if line_s.startswith("# "):
            p = doc.add_paragraph()
            set_p_rtl(p, space_before=14, space_after=8, align="center")
            add_r(p, line_s[2:], size_pt=16, bold=True, color_rgb=(0x1A, 0x23, 0x7E))
        elif line_s.startswith("## "):
            p = doc.add_paragraph()
            set_p_rtl(p, space_before=12, space_after=6, align="right")
            add_r(p, line_s[3:], size_pt=13, bold=True, color_rgb=(0x0D, 0x47, 0xA1))
        elif line_s.startswith("### "):
            p = doc.add_paragraph()
            set_p_rtl(p, space_before=8, space_after=4, align="right")
            add_r(p, line_s[4:], size_pt=11.5, bold=True, color_rgb=(0x15, 0x65, 0xC0))
        elif line_s.startswith("- ") or line_s.startswith("* "):
            p = doc.add_paragraph()
            set_p_rtl(p, space_before=2, space_after=2, align="both")
            add_r(p, "• " + line_s[2:], size_pt=10.5)
        elif line_s.startswith("> "):
            p = doc.add_paragraph()
            set_p_rtl(p, space_before=4, space_after=4, align="both")
            add_r(p, line_s[2:], size_pt=10, italic=True, color_rgb=(0x42, 0x42, 0x42))
        else:
            p = doc.add_paragraph()
            set_p_rtl(p, space_before=2, space_after=3.5, align="both")
            add_r(p, line_s, size_pt=11)

    out_path = "MEDICAL_PROPOSAL_LUPEOL_NDV_V87_REAL_WORLD.docx"
    doc.save(out_path)
    print(f"Saved audited Word proposal successfully to: {out_path}")

if __name__ == "__main__":
    build_audited_docx()
